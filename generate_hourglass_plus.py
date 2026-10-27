import sys
import os
import csv
import html
import re
import hashlib
import openpyxl
from openpyxl.styles import Font, PatternFill

SECRET_SALT = "McConigleyReunion2026"

def get_obscure_slug(prefix):
    token = hashlib.sha256(f"{prefix.lower()}-{SECRET_SALT}".encode()).hexdigest()[:8]
    return f"{prefix.lower()}-{token}"

def parse_gedcom(gedcom_path):
    individuals = {}
    families = {}
    current = None
    target = None
    context = None

    with open(gedcom_path, "r", encoding="utf-8-sig", errors="ignore") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            parts = line.split(" ", 2)
            lvl = parts[0]
            tag = parts[1] if len(parts) > 1 else ""
            val = parts[2] if len(parts) > 2 else ""

            if lvl == "0":
                if val == "INDI":
                    current = {
                        "id": tag.strip("@"), "name": "", "alt_names": [],
                        "birth_date": "", "birth_place": "",
                        "death_date": "", "death_place": "",
                        "fams": [], "famc": None
                    }
                    individuals[tag] = current
                    target = "INDI"
                    context = None
                elif val == "FAM":
                    current = {
                        "id": tag.strip("@"), "husb": None, "wife": None,
                        "children": [], "marr_date": "", "marr_place": ""
                    }
                    families[tag] = current
                    target = "FAM"
                    context = None
                else:
                    target = None
                    context = None

            elif target == "INDI":
                if lvl == "1":
                    context = tag
                    if tag == "NAME":
                        clean_n = val.replace("/", "").strip()
                        if not current["name"]:
                            current["name"] = clean_n
                        else:
                            current["alt_names"].append(clean_n)
                    elif tag == "FAMS":
                        current["fams"].append(val)
                    elif tag == "FAMC":
                        current["famc"] = val
                elif lvl == "2":
                    if context == "BIRT":
                        if tag == "DATE": current["birth_date"] = val
                        elif tag == "PLAC": current["birth_place"] = val
                    elif context == "DEAT":
                        if tag == "DATE": current["death_date"] = val
                        elif tag == "PLAC": current["death_place"] = val

            elif target == "FAM":
                if lvl == "1":
                    context = tag
                    if tag == "HUSB": current["husb"] = val
                    elif tag == "WIFE": current["wife"] = val
                    elif tag == "CHIL": current["children"].append(val)
                elif lvl == "2" and context == "MARR":
                    if tag == "DATE": current["marr_date"] = val
                    elif tag == "PLAC": current["marr_place"] = val

    return individuals, families

def find_target_person(individuals, query):
    q = query.strip()
    clean_id = q.strip("@")
    for pid, data in individuals.items():
        if data["id"] == clean_id or pid == f"@{clean_id}@":
            return pid

    year_match = re.search(r"\b(1\d{3}|20\d{2})\b", q)
    target_year = year_match.group(1) if year_match else None
    name_fragment = re.sub(r"\b(1\d{3}|20\d{2})\b", "", q).strip().lower()

    matches = []
    for pid, data in individuals.items():
        all_names = [data["name"]] + data.get("alt_names", [])
        if any(name_fragment in n.lower() for n in all_names):
            bdate = data.get("birth_date", "")
            if target_year:
                if target_year in bdate:
                    matches.append((pid, data))
            else:
                matches.append((pid, data))

    if not matches:
        return None
    if len(matches) == 1:
        return matches[0][0]

    print(f"\nMultiple matches found for '{query}':")
    for idx, (pid, d) in enumerate(matches, 1):
        b = d.get("birth_date") or "Unknown"
        d_val = d.get("death_date") or "Living/Unknown"
        print(f"  [{idx}] {d['name']} (b. {b} - d. {d_val}) [ID: {d['id']}]")
    
    choice = input("\nEnter selection number (or Enter for [1]): ").strip()
    if choice.isdigit() and 1 <= int(choice) <= len(matches):
        return matches[int(choice) - 1][0]
    return matches[0][0]

def build_hourglass_plus(gedcom_path, focus_query, base_name):
    if not os.path.exists(gedcom_path):
        print(f"Error: GEDCOM file '{gedcom_path}' not found.")
        return

    clean_base = os.path.splitext(base_name)[0]
    slug = get_obscure_slug(clean_base)

    os.makedirs(os.path.join("family", "sheets"), exist_ok=True)

    out_csv = os.path.join("family", "sheets", f"{slug}.csv")
    out_xlsx = os.path.join("family", "sheets", f"{slug}.xlsx")
    out_html = os.path.join("family", f"{slug}.html")
    sheet_redirect_link = f"sheets/redirect.php?id={slug}"

    indis, fams = parse_gedcom(gedcom_path)
    if len(indis) == 0:
        print(f"\n❌ Error: '{gedcom_path}' contains 0 people! Check your Gramps export file.\n")
        return

    focus_id = find_target_person(indis, focus_query)
    if not focus_id:
        print(f"Error: Could not locate focus person from '{focus_query}'.")
        return

    focus_p = indis[focus_id]
    output_rows = []

    def get_marriage_info(p1_id, p2_id):
        for f_id in indis.get(p1_id, {}).get("fams", []):
            if f_id in fams:
                fam = fams[f_id]
                if (fam["husb"] == p1_id and fam["wife"] == p2_id) or (fam["husb"] == p2_id and fam["wife"] == p1_id):
                    return fam["marr_date"].strip(), fam["marr_place"].strip()
        return "", ""

    def get_max_depth(pid, depth=0):
        famc_id = indis[pid].get("famc")
        if not famc_id or famc_id not in fams:
            return depth
        fam = fams[famc_id]
        f_depth = get_max_depth(fam["husb"], depth + 1) if fam["husb"] and fam["husb"] in indis else depth + 1
        m_depth = get_max_depth(fam["wife"], depth + 1) if fam["wife"] and fam["wife"] in indis else depth + 1
        return max(f_depth, m_depth)

    max_anc_depth = max(1, get_max_depth(focus_id))
    focus_col_idx = max_anc_depth

    # 1. ANCESTORS
    output_rows.append({"type": "section", "title": "Direct Ancestral Line"})
    visited_anc_couples = set()

    def add_ancestor_branch(pid, child_gen_from_focus):
        famc_id = indis[pid].get("famc")
        if not famc_id or famc_id not in fams:
            return

        fam = fams[famc_id]
        f_id = fam["husb"]
        m_id = fam["wife"]
        couple_key = (f_id, m_id)

        if couple_key in visited_anc_couples:
            return

        if f_id and f_id in indis:
            add_ancestor_branch(f_id, child_gen_from_focus + 1)
        if m_id and m_id in indis:
            add_ancestor_branch(m_id, child_gen_from_focus + 1)

        visited_anc_couples.add(couple_key)
        gen = child_gen_from_focus
        col_slot = focus_col_idx - gen

        if gen == 1: f_role, m_role = "Father", "Mother"
        elif gen == 2: f_role, m_role = "Grandfather", "Grandmother"
        elif gen == 3: f_role, m_role = "Great-Grandfather", "Great-Grandmother"
        else: f_role, m_role = f"Ancestor M (-{gen})", f"Ancestor F (-{gen})"

        marr_date, marr_place = get_marriage_info(f_id, m_id) if f_id and m_id else ("", "")

        if f_id and f_id in indis:
            fp = indis[f_id]
            output_rows.append({
                "type": "ancestor",
                "gramps_id": fp["id"],
                "level_label": f_role,
                "col_idx": col_slot,
                "name": fp["name"],
                "birth_date": fp["birth_date"].strip(),
                "birth_place": fp["birth_place"].strip(),
                "death_date": fp["death_date"].strip(),
                "death_place": fp["death_place"].strip(),
                "marr_date": marr_date,
                "marr_place": marr_place
            })

        if m_id and m_id in indis:
            mp = indis[m_id]
            output_rows.append({
                "type": "spouse",
                "gramps_id": mp["id"],
                "level_label": f"sp. ({m_role})",
                "col_idx": col_slot,
                "name": f"sp. {mp['name']}",
                "birth_date": mp["birth_date"].strip(),
                "birth_place": mp["birth_place"].strip(),
                "death_date": mp["death_date"].strip(),
                "death_place": mp["death_place"].strip(),
                "marr_date": marr_date,
                "marr_place": marr_place
            })

    add_ancestor_branch(focus_id, 1)

    # 2. FOCUS & SIBLINGS
    output_rows.append({"type": "section", "title": "Focus Generation & Siblings"})

    parent_famc = focus_p.get("famc")
    sibling_ids = []
    if parent_famc and parent_famc in fams:
        sibling_ids = [c for c in fams[parent_famc]["children"] if c != focus_id]

    for sib_id in sibling_ids:
        sib = indis[sib_id]
        output_rows.append({
            "type": "sibling",
            "gramps_id": sib["id"],
            "level_label": "Sib.",
            "col_idx": focus_col_idx,
            "name": sib["name"],
            "birth_date": sib["birth_date"].strip(),
            "birth_place": sib["birth_place"].strip(),
            "death_date": sib["death_date"].strip(),
            "death_place": sib["death_place"].strip(),
            "marr_date": "",
            "marr_place": ""
        })
        for f_id in sib["fams"]:
            if f_id in fams:
                fam = fams[f_id]
                sp_id = fam["wife"] if fam["husb"] == sib_id else fam["husb"]
                if sp_id and sp_id in indis:
                    sp = indis[sp_id]
                    m_date, m_place = fam["marr_date"].strip(), fam["marr_place"].strip()
                    output_rows.append({
                        "type": "spouse",
                        "gramps_id": sp["id"],
                        "level_label": "sp.",
                        "col_idx": focus_col_idx,
                        "name": f"sp. {sp['name']}",
                        "birth_date": sp["birth_date"].strip(),
                        "birth_place": sp["birth_place"].strip(),
                        "death_date": sp["death_date"].strip(),
                        "death_place": sp["death_place"].strip(),
                        "marr_date": m_date,
                        "marr_place": m_place
                    })

    output_rows.append({
        "type": "focus",
        "gramps_id": focus_p["id"],
        "level_label": "1. FOCUS",
        "col_idx": focus_col_idx,
        "name": focus_p["name"],
        "birth_date": focus_p["birth_date"].strip(),
        "birth_place": focus_p["birth_place"].strip(),
        "death_date": focus_p["death_date"].strip(),
        "death_place": focus_p["death_place"].strip(),
        "marr_date": "",
        "marr_place": ""
    })

    for f_id in focus_p["fams"]:
        if f_id in fams:
            fam = fams[f_id]
            sp_id = fam["wife"] if fam["husb"] == focus_id else fam["husb"]
            if sp_id and sp_id in indis:
                sp = indis[sp_id]
                m_date, m_place = fam["marr_date"].strip(), fam["marr_place"].strip()
                output_rows.append({
                    "type": "spouse",
                    "gramps_id": sp["id"],
                    "level_label": "sp.",
                    "col_idx": focus_col_idx,
                    "name": f"sp. {sp['name']}",
                    "birth_date": sp["birth_date"].strip(),
                    "birth_place": sp["birth_place"].strip(),
                    "death_date": sp["death_date"].strip(),
                    "death_place": sp["death_place"].strip(),
                    "marr_date": m_date,
                    "marr_place": m_place
                })

    # 3. DESCENDANTS
    output_rows.append({"type": "section", "title": "Descendant Generations"})
    visited_desc = set()

    def walk_descendants(pid, gen_num=2):
        if pid in visited_desc or pid not in indis: return
        visited_desc.add(pid)
        p = indis[pid]

        current_col = focus_col_idx + (gen_num - 1)

        output_rows.append({
            "type": "descendant",
            "gramps_id": p["id"],
            "level_label": f"{gen_num}.",
            "col_idx": current_col,
            "name": f"{gen_num}. {p['name']}",
            "birth_date": p["birth_date"].strip(),
            "birth_place": p["birth_place"].strip(),
            "death_date": p["death_date"].strip(),
            "death_place": p["death_place"].strip(),
            "marr_date": "",
            "marr_place": ""
        })

        for f_id in p["fams"]:
            if f_id in fams:
                fam = fams[f_id]
                sp_id = fam["wife"] if fam["husb"] == pid else fam["husb"]
                if sp_id and sp_id in indis:
                    sp = indis[sp_id]
                    m_date, m_place = fam["marr_date"].strip(), fam["marr_place"].strip()
                    output_rows.append({
                        "type": "spouse",
                        "gramps_id": sp["id"],
                        "level_label": "sp.",
                        "col_idx": current_col,
                        "name": f"sp. {sp['name']}",
                        "birth_date": sp["birth_date"].strip(),
                        "birth_place": sp["birth_place"].strip(),
                        "death_date": sp["death_date"].strip(),
                        "death_place": sp["death_place"].strip(),
                        "marr_date": m_date,
                        "marr_place": m_place
                    })
                for child_id in fam["children"]:
                    walk_descendants(child_id, gen_num + 1)

    for f_id in focus_p["fams"]:
        if f_id in fams:
            for child_id in fams[f_id]["children"]:
                walk_descendants(child_id, 2)

    max_col = max(r.get("col_idx", 0) for r in output_rows if r["type"] != "section")
    total_cols = max_col + 1

    col_headers = [""] * total_cols
    for c in range(total_cols):
        if c < focus_col_idx:
            dist = focus_col_idx - c
            if dist == 1: col_headers[c] = "Parents"
            elif dist == 2: col_headers[c] = "Grandparents"
            elif dist == 3: col_headers[c] = "Great-Grands"
            else: col_headers[c] = f"Ancestors (-{dist})"
        elif c == focus_col_idx:
            col_headers[c] = "Focus / Siblings"
        else:
            g_num = c - focus_col_idx + 1
            if g_num == 2: col_headers[c] = "Children (Gen 2)"
            elif g_num == 3: col_headers[c] = "Grandchildren (Gen 3)"
            elif g_num == 4: col_headers[c] = "Great-Grands (Gen 4)"
            elif g_num == 5: col_headers[c] = "Great x2 (Gen 5)"
            else: col_headers[c] = f"Descendants (Gen {g_num})"

    full_headers = (
        ["Gramps ID", "Role / Line"]
        + col_headers
        + ["Birth Date", "Birth Place", "Death Date", "Death Place", "Marriage Date", "Marriage Place", "Corrections", "Notes"]
    )

    with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(full_headers)
        for r in output_rows:
            if r["type"] == "section":
                writer.writerow(["", f"=== {r['title'].upper()} ==="] + [""] * (len(full_headers) - 2))
            else:
                row_cells = [""] * total_cols
                row_cells[r["col_idx"]] = r["name"]
                writer.writerow(
                    [r["gramps_id"], r["level_label"]]
                    + row_cells
                    + [r["birth_date"], r["birth_place"], r["death_date"], r["death_place"], r["marr_date"], r["marr_place"], "", ""]
                )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = clean_base
    ws.append(full_headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")

    for r in output_rows:
        if r["type"] == "section":
            row_vals = ["", f"=== {r['title'].upper()} ==="] + [""] * (len(full_headers) - 2)
            ws.append(row_vals)
            cur_row = ws.max_row
            ws.cell(cur_row, 2).font = Font(bold=True, color="FFFFFF")
            ws.cell(cur_row, 2).fill = PatternFill("solid", fgColor="334155")
        else:
            row_cells = [""] * total_cols
            row_cells[r["col_idx"]] = r["name"]
            row_vals = (
                [r["gramps_id"], r["level_label"]]
                + row_cells
                + [r["birth_date"], r["birth_place"], r["death_date"], r["death_place"], r["marr_date"], r["marr_place"], "", ""]
            )
            ws.append(row_vals)
            cur_row = ws.max_row
            if r["type"] == "focus":
                for col_i in range(1, len(row_vals) + 1):
                    ws.cell(cur_row, col_i).fill = PatternFill("solid", fgColor="EFF6FF")
            elif r["type"] == "spouse":
                for col_i in range(1, len(row_vals) + 1):
                    ws.cell(cur_row, col_i).fill = PatternFill("solid", fgColor="FFFDF5")

    wb.save(out_xlsx)
    print(f"Generated Excel: {out_xlsx}")

    table_rows = ""
    for r in output_rows:
        if r["type"] == "section":
            table_rows += f'<tr class="section-divider"><td colspan="5">{html.escape(r["title"])}</td></tr>'
        else:
            cls = f"row-{r['type']}"
            indent_px = r["col_idx"] * 22 + 8
            tree_guide = '<span class="tree-guide">└─ </span>' if r["col_idx"] > 0 else ''
            table_rows += f"""
            <tr class="{cls}">
              <td class="badge-cell"><span class="badge {r['type']}">{html.escape(r['level_label'])}</span></td>
              <td class="name-cell" style="padding-left: {indent_px}px;">
                {tree_guide}<strong>{html.escape(r['name'])}</strong>
              </td>
              <td class="date-cell">{html.escape(r['birth_date'])}</td>
              <td class="date-cell">{html.escape(r['death_date'])}</td>
              <td class="date-cell">{html.escape(r['marr_date'])}</td>
            </tr>"""

    html_doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <meta charset="utf-8">
  <title>Hourglass+ View: {html.escape(focus_p['name'])}</title>
  <style>
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
      margin: 24px;
      background: #f8fafc;
      color: #0f172a;
    }}
    .header-bar {{
      display: flex;
      justify-content: space-between;
      align-items: center;
      max-width: fit-content;
      min-width: 680px;
      margin-bottom: 16px;
    }}
    h2 {{ margin: 0 0 4px 0; }}
    .subtitle {{ color: #64748b; margin: 0; font-size: 0.9rem; }}
    .edit-btn {{
      display: inline-block;
      background: #0284c7;
      color: #fff;
      text-decoration: none;
      font-size: 0.82rem;
      font-weight: 600;
      padding: 8px 14px;
      border-radius: 6px;
      box-shadow: 0 1px 2px rgba(0,0,0,0.05);
      transition: background 0.15s ease-in-out;
    }}
    .edit-btn:hover {{ background: #0369a1; }}
    table {{
      width: fit-content;
      min-width: 680px;
      max-width: 100%;
      border-collapse: collapse;
      background: #fff;
      box-shadow: 0 1px 3px rgba(0,0,0,0.08);
      border-radius: 6px;
      overflow: hidden;
    }}
    th, td {{ padding: 6px 12px; font-size: 0.86rem; border-bottom: 1px solid #f1f5f9; }}
    th {{
      background: #0f172a;
      color: #f8fafc;
      text-align: left;
      font-size: 0.76rem;
      text-transform: uppercase;
      letter-spacing: 0.5px;
    }}
    th.badge-header, td.badge-cell {{ width: 1%; white-space: nowrap; padding-right: 6px; }}
    th.name-header, td.name-cell {{ white-space: nowrap; padding-right: 24px; }}
    th.date-header, td.date-cell {{
      width: 125px;
      min-width: 110px;
      white-space: nowrap;
      font-variant-numeric: tabular-nums;
      color: #334155;
      padding-left: 12px;
      padding-right: 12px;
    }}
    .section-divider td {{
      background: #334155;
      color: #f8fafc;
      font-weight: 700;
      text-transform: uppercase;
      letter-spacing: 0.8px;
      font-size: 0.76rem;
      padding: 8px 12px;
    }}
    .row-focus {{ background: #eff6ff; }}
    .row-focus strong {{ color: #1d4ed8; }}
    .row-spouse {{ background: #fffdf5; color: #78350f; }}
    .row-ancestor {{ background: #fdfefe; }}
    .row-descendant:hover, .row-ancestor:hover, .row-sibling:hover {{ background: #f1f5f9; }}
    .tree-guide {{ color: #94a3b8; font-family: monospace; font-weight: bold; margin-right: 4px; }}
    .badge {{ display: inline-block; padding: 2px 6px; border-radius: 4px; font-size: 0.70rem; font-weight: 600; }}
    .badge.ancestor {{ background: #e0e7ff; color: #3730a3; }}
    .badge.focus {{ background: #2563eb; color: #fff; }}
    .badge.spouse {{ background: #fef3c7; color: #92400e; }}
    .badge.sibling {{ background: #f1f5f9; color: #475569; }}
    .badge.descendant {{ background: #dcfce7; color: #166534; }}
    .name-cell {{ border-left: 2px solid #e2e8f0; }}
  </style>
</head>
<body>
  <div class="header-bar">
    <div>
      <h2>🌳 Extended Hourglass: {html.escape(focus_p['name'])}</h2>
      <div class="subtitle">Direct Ancestors • Siblings • Focus Person &amp; Spouse • Descendants</div>
    </div>
    <a href="{html.escape(sheet_redirect_link)}" target="_blank" class="edit-btn">✏️ Report changes</a>
  </div>
  <div style="overflow-x: auto; -webkit-overflow-scrolling: touch;">
    <table>
    <thead>
      <tr>
        <th class="badge-header">Role / Line</th>
        <th class="name-header">Individual &amp; Generation Link</th>
        <th class="date-header">Birth Date</th>
        <th class="date-header">Death Date</th>
        <th class="date-header">Marriage Date</th>
      </tr>
    </thead>
    <tbody>{table_rows}</tbody>
      </table>
</div>
</body>
</html>"""

    with open(out_html, "w", encoding="utf-8") as f:
        f.write(html_doc)
    print(f"Generated HTML:  {out_html}")

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python3 generate_hourglass_plus.py <tree.ged> <Focus Query> <Output Base Name>")
        sys.exit(1)

    ged_file = sys.argv[1]
    focus_target = sys.argv[2]
    base_slug = sys.argv[3]
    build_hourglass_plus(ged_file, focus_target, base_slug)
